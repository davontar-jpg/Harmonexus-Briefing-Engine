from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

LEGACY_SHEETS = {"Dashboard", "Signal_Engine", "Machine_Input", "Briefing"}
V5_SHEETS = {"Calculated_Signals", "Instrument_Scores", "Score_History", "System_Log"}


def validate(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    names = set(workbook.sheetnames)
    return {
        "path": str(path),
        "sheet_count": len(names),
        "legacy_compatible": LEGACY_SHEETS.issubset(names),
        "missing_legacy_sheets": sorted(LEGACY_SHEETS - names),
        "v5_installed": V5_SHEETS.issubset(names),
        "missing_v5_sheets": sorted(V5_SHEETS - names),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Read-only Harmonexus workbook contract validation")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    result = validate(args.workbook)
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result["legacy_compatible"] else 1)
